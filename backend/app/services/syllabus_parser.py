"""Extract clean text from uploaded syllabus files (PDF, DOCX, TXT). Works on in-memory bytes,
so nothing is written to disk."""
import io
import logging
import re

from app.core.config import settings
from app.core.exceptions import AppError, FileValidationError
from app.utils.validators import validate_upload

logger = logging.getLogger(__name__)


def clean_text(text: str) -> str:
    text = text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t\u00a0]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_text_from_pdf(data: bytes) -> str:
    try:
        import fitz  # PyMuPDF
    except ImportError as exc:  # pragma: no cover
        raise AppError("PDF support is not installed on the server (PyMuPDF).", status_code=500) from exc

    try:
        doc = fitz.open(stream=data, filetype="pdf")
    except Exception as exc:
        logger.warning("PDF open failed: %s", exc)
        raise FileValidationError("The PDF appears to be corrupted and could not be opened.") from exc

    try:
        if doc.needs_pass:
            raise FileValidationError("The PDF is password-protected. Please upload an unlocked copy.")
        pages = [page.get_text("text") for page in doc]
        page_count = len(pages)
    except FileValidationError:
        raise
    except Exception as exc:
        logger.warning("PDF read failed: %s", exc)
        raise FileValidationError("The PDF could not be read; it may be corrupted.") from exc
    finally:
        doc.close()

    text = clean_text("\n\n".join(pages))
    if not text:
        raise FileValidationError(
            f"No extractable text found in this PDF ({page_count} page(s)). It is probably a scanned or "
            "image-only document, which requires OCR. Please upload a text-based PDF, a DOCX, or a TXT file."
        )
    return text


def extract_text_from_docx(data: bytes) -> str:
    from docx import Document

    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:
        logger.warning("DOCX open failed: %s", exc)
        raise FileValidationError("The DOCX file appears to be corrupted and could not be opened.") from exc

    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    text = clean_text("\n".join(parts))
    if not text:
        raise FileValidationError("The DOCX document contains no text.")
    return text


def extract_text_from_txt(data: bytes) -> str:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("latin-1")
    text = clean_text(text)
    if not text:
        raise FileValidationError("The text file contains no text.")
    return text


def extract_text_from_file(filename: str | None, data: bytes) -> tuple[str, str]:
    """Validate + extract. Returns (text, extension_without_dot)."""
    ext = validate_upload(filename, data, settings.max_upload_bytes)
    if ext == ".pdf":
        text = extract_text_from_pdf(data)
    elif ext == ".docx":
        text = extract_text_from_docx(data)
    else:
        text = extract_text_from_txt(data)
    if len(text) < settings.min_syllabus_chars:
        raise FileValidationError("The syllabus text is too short to analyse. Is this the right document?")
    return text, ext.lstrip(".")
